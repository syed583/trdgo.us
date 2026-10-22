"""Data Health Check for the Earnings & News tabs (read-only; no model logic touched).

For every data parameter the earnings models (Claude, Grok, Meta, ChatGPT, Kimi,
Gemini, Claude Pro, Pre-Earnings Direction, IV Crush Check, Unusual Activity) and
News/Calendar tabs read, this calls the real provider endpoint directly (so the
HTTP code is the provider's own) on AAPL, MSFT, NVDA, AMZN, TSLA plus the two
largest names reporting earnings this week, then writes one row per parameter to
the "Data Health" sheet of Data_Health_Check.xlsx.

    NO DATA : response empty, null or zero for every ticker tested
    ERROR   : a request failed (HTTP >= 400 other than 404, or network error)
    OK      : at least one ticker returned real data and nothing errored

Run:  python data_health_check.py [--out path.xlsx]
Keys are read from backend/.env (UNUSUAL_WHALES_API_KEY, FINNHUB_API_KEY) and are
never printed or written to the workbook.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, timedelta
from pathlib import Path

CORE = ["AAPL", "MSFT", "NVDA", "AMZN", "TSLA"]
UW_BASE = "https://api.unusualwhales.com"
FH_BASE = "https://finnhub.io/api/v1"
HERE = Path(__file__).resolve().parent


def _load_env() -> dict:
    env = dict(os.environ)
    p = HERE / ".env"
    if p.exists():
        for line in p.read_text(encoding="utf-8", errors="ignore").splitlines():
            if "=" in line and not line.lstrip().startswith("#"):
                k, v = line.split("=", 1)
                env.setdefault(k.strip(), v.strip().strip("\"'"))
    return env


ENV = _load_env()
UW_KEY = ENV.get("UNUSUAL_WHALES_API_KEY", "")
FH_KEY = ENV.get("FINNHUB_API_KEY", "")


# --------------------------------------------------------------------- HTTP --
_last_fh = 0.0


def _http(provider: str, path: str, params: dict | None = None) -> dict:
    """-> {code, data, error}. Retries once on 429/5xx. Never logs the key."""
    global _last_fh
    if provider == "UW":
        if not UW_KEY:
            return {"code": None, "data": None, "error": "UNUSUAL_WHALES_API_KEY not set"}
        base, hdr = UW_BASE, {"Authorization": f"Bearer {UW_KEY}", "Accept": "application/json",
                              "User-Agent": "TradeGo-DataHealth/1.0"}
    else:
        if not FH_KEY:
            return {"code": None, "data": None, "error": "FINNHUB_API_KEY not set"}
        base, hdr = FH_BASE, {"X-Finnhub-Token": FH_KEY, "User-Agent": "TradeGo-DataHealth/1.0"}
        gap = 1.2 - (time.time() - _last_fh)          # stay under 55 calls/min
        if gap > 0:
            time.sleep(gap)
    q = urllib.parse.urlencode({k: v for k, v in (params or {}).items() if v not in (None, "")})
    url = base + path + (("?" + q) if q else "")
    out = {"code": None, "data": None, "error": ""}
    for attempt in range(2):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=hdr), timeout=30) as r:
                out["code"] = r.status
                body = r.read().decode("utf-8", "ignore") or "null"
                out["data"] = json.loads(body)
                out["error"] = ""
                break
        except urllib.error.HTTPError as e:
            out["code"] = e.code
            try:
                b = json.loads(e.read().decode("utf-8", "ignore") or "{}")
                out["error"] = str(b.get("message") or b.get("error") or b)[:200]
            except Exception:  # noqa: BLE001
                out["error"] = f"HTTP {e.code}"
            if e.code in (429, 500, 502, 503, 504) and attempt == 0:
                time.sleep(2.0)
                continue
            break
        except Exception as e:  # noqa: BLE001
            out["error"] = f"{type(e).__name__}: {e}"[:200]
            if attempt == 0:
                time.sleep(1.0)
                continue
            break
    if provider != "UW":
        _last_fh = time.time()
    else:
        time.sleep(0.12)
    return out


# -------------------------------------------------------------- emptiness ----
def _leaves(x, depth=0):
    if depth > 6:
        return
    if isinstance(x, dict):
        for v in x.values():
            yield from _leaves(v, depth + 1)
    elif isinstance(x, list):
        for v in x:
            yield from _leaves(v, depth + 1)
    else:
        yield x


def _is_zero(v) -> bool:
    if v is None or v is False or v == "":
        return True
    try:
        return float(v) == 0.0
    except (TypeError, ValueError):
        return False                                   # a real string/date is data


def is_empty(data) -> bool:
    """Empty / null / all-zero payload."""
    if data is None:
        return True
    if isinstance(data, dict) and "data" in data:      # UW wraps the answer in "data"
        data = data["data"]
    return all(_is_zero(v) for v in _leaves(data))


def sample(data) -> str:
    if isinstance(data, dict) and "data" in data:
        data = data["data"]
    if isinstance(data, list):
        head = data[0] if data else None
        s = f"{len(data)} rows; " + json.dumps(head, default=str)
    else:
        s = json.dumps(data, default=str)
    return s[:110]


# ----------------------------------------------------------------- catalog ---
# (parameter, used by, provider, endpoint template, params, scope)
#   scope "t" = run per ticker ({t} substituted); "m" = market-wide, run once.
EARN_TABS = "Claude, Grok, Meta, ChatGPT, Kimi, Gemini, Claude Pro, Pre-Earnings Direction"
THIS_WEEK = "__week__"
TODAY = date.today().isoformat()

CATALOG = [
    ("EPS estimates & revisions", EARN_TABS, "UW", "/api/companies/{t}/earnings-estimates", {}, "t"),
    ("Revenue estimates & revisions", EARN_TABS, "UW", "/api/companies/{t}/earnings-estimates", {}, "t"),
    ("Forward guidance / beat-raise", EARN_TABS, "UW", "/api/companies/{t}/earnings-estimates", {}, "t"),
    ("Historical earnings reaction", "Claude, Grok, Meta, ChatGPT, Kimi, Claude Pro", "UW", "/api/earnings/{t}", {"limit": 24}, "t"),
    ("Price action / relative strength (daily OHLC)", EARN_TABS, "UW", "/api/stock/{t}/ohlc/1d", {"limit": 60}, "t"),
    ("Live quote / prior close", EARN_TABS + ", IV Crush Check", "UW", "/api/stock/{t}/quote", {}, "t"),
    ("Sector & market regime (SPY quote)", "Claude, Grok, Meta, ChatGPT, Kimi, Gemini, Claude Pro", "UW", "/api/stock/SPY/quote", {}, "m"),
    ("VIX index (raw endpoint)", "Kimi, Claude Pro", "UW", "/api/stock/VIX/quote", {}, "m"),
    ("VIX regime proxy (VIXY ETF quote)", "Kimi, Claude Pro", "UW", "/api/stock/VIXY/quote", {}, "m"),
    ("Market tide (market flow)", "Pre-Earnings Direction, Unusual Activity", "UW", "/api/market/market-tide", {}, "m"),
    ("Options flow / sweeps", EARN_TABS, "UW", "/api/stock/{t}/flow-alerts", {"limit": 50}, "t"),
    ("Unusual options activity", EARN_TABS + ", Unusual Activity", "UW", "/api/option-activity/unusual",
     {"ticker_symbol": "{t}", "min_premium": 50000, "limit": 50, "unusual": "true"}, "t"),
    ("OI positioning / OI change", EARN_TABS, "UW", "/api/stock/{t}/oi-change", {"limit": 100}, "t"),
    ("Options volume / put-call", "Claude, Grok, Meta, ChatGPT, Kimi, Gemini, Unusual Activity", "UW", "/api/stock/{t}/options-volume", {}, "t"),
    ("Net premium ticks (call/put premium)", "Claude, Meta, Unusual Activity", "UW", "/api/stock/{t}/net-prem-ticks", {}, "t"),
    ("Disparity / skew (option contracts)", EARN_TABS + ", Unusual Activity", "UW", "/api/stock/{t}/option-contracts", {"limit": 200}, "t"),
    ("IV rank / implied volatility", "IV Crush Check, Pre-Earnings Direction, Meta, Kimi, Gemini", "UW", "/api/stock/{t}/iv-rank", {}, "t"),
    ("Max pain / key option levels", "Pre-Earnings Direction, IV Crush Check", "UW", "/api/stock/{t}/max-pain", {}, "t"),
    ("Dark pool prints", EARN_TABS, "UW", "/api/darkpool/{t}", {"limit": 100}, "t"),
    ("Volume price levels (dark-pool levels)", "Claude, Grok, Meta, ChatGPT", "UW", "/api/stock/{t}/stock-volume-price-levels", {}, "t"),
    ("Short interest & days-to-cover", EARN_TABS, "UW", "/api/shorts/{t}/interest-float/v2", {}, "t"),
    ("Insider transactions (open-market buys)", "Claude, Grok, Meta, ChatGPT, Kimi, Gemini, Claude Pro", "UW", "/api/insider/transactions",
     {"ticker_symbol": "{t}", "limit": 100}, "t"),
    ("Insider flow summary", "Claude, Grok, Meta, ChatGPT, Kimi", "UW", "/api/insider/{t}", {}, "t"),
    ("Institutional ownership / big funds (13F)", "Claude, Claude Pro, Pre-Earnings Direction", "UW", "/api/institution/{t}/ownership", {"limit": 100}, "t"),
    ("Analyst ratings (scored: Grok, ChatGPT, Kimi, Claude Pro)", "Grok, ChatGPT, Kimi, Claude Pro", "UW", "/api/screener/analysts", {"ticker": "{t}", "limit": 20}, "t"),
    ("News headlines", "News & Sentiment", "UW", "/api/news/headlines", {"ticker": "{t}", "limit": 20}, "t"),
    ("Earnings calendar / timing / implied move (pre-market)", "Earnings Calendar, IV Crush Check, all models", "UW", "/api/earnings/premarket", {"date": THIS_WEEK, "limit": 500}, "m"),
    ("Earnings calendar / timing / implied move (after-hours)", "Earnings Calendar, IV Crush Check, all models", "UW", "/api/earnings/afterhours", {"date": THIS_WEEK, "limit": 500}, "m"),
    ("Peer read-through: peers list", "Claude, Grok, Meta, ChatGPT, Kimi, Gemini, Claude Pro", "FH", "/stock/peers", {"symbol": "{t}"}, "t"),
    ("Analyst recommendations (Finnhub)", "Grok, Kimi, Claude Pro, ChatGPT", "FH", "/stock/recommendation", {"symbol": "{t}"}, "t"),
    ("Guidance hit-rate: 8Q EPS beat history (Finnhub)", "Gemini", "FH", "/stock/earnings", {"symbol": "{t}"}, "t"),
    ("Fundamental metrics (Finnhub)", "Grok, Claude Pro", "FH", "/stock/metric", {"symbol": "{t}", "metric": "all"}, "t"),
    ("Insider sentiment (Finnhub)", "Claude, Grok, Meta, ChatGPT", "FH", "/stock/insider-sentiment",
     {"symbol": "{t}", "from": (date.today() - timedelta(days=180)).isoformat(), "to": TODAY}, "t"),
    ("Earnings calendar (Finnhub)", "Earnings Calendar, Claude, Grok", "FH", "/calendar/earnings",
     {"symbol": "{t}", "from": TODAY, "to": (date.today() + timedelta(days=90)).isoformat()}, "t"),
    ("Company news (Finnhub)", "News & Sentiment", "FH", "/company-news",
     {"symbol": "{t}", "from": (date.today() - timedelta(days=7)).isoformat(), "to": TODAY}, "t"),
]

# Model parameters that have no data source wired at all.
UNWIRED: list[tuple] = []     # every scored parameter now has a wired source


# ------------------------------------------------------------------ driver ---
def pick_week_tickers(exclude: list[str]) -> list[str]:
    """Two largest-market-cap names reporting today..Friday (UW calendar)."""
    today = date.today()
    days = [today + timedelta(days=i) for i in range(0, 7 - today.weekday()) if (today + timedelta(days=i)).weekday() < 5]
    best: dict[str, float] = {}
    for d in days:
        for path in ("/api/earnings/premarket", "/api/earnings/afterhours"):
            r = _http("UW", path, {"date": d.isoformat(), "limit": 500})
            rows = r["data"].get("data") if isinstance(r["data"], dict) else None
            for row in rows or []:
                s = (row.get("symbol") or "").upper()
                try:
                    cap = float(row.get("marketcap") or 0)
                except ValueError:
                    cap = 0
                if s and s not in exclude and row.get("has_options") and cap:
                    best[s] = max(best.get(s, 0), cap)
    return [s for s, _ in sorted(best.items(), key=lambda kv: -kv[1])[:2]]


def _subst(v, t):
    if isinstance(v, str):
        return v.replace("{t}", t)
    return v


def run(out_path: Path) -> list[dict]:
    week = pick_week_tickers(CORE)
    tickers = CORE + week
    print(f"Tickers: {', '.join(tickers)}  (this week: {', '.join(week) or 'none found'})")
    cache: dict[tuple, dict] = {}
    rows: list[dict] = []
    for name, used, prov, path, params, scope in CATALOG:
        targets = tickers if scope == "t" else [""]
        calls = []
        for t in targets:
            p = {k: (_subst(v, t)) for k, v in params.items()}
            if p.get("date") == THIS_WEEK:
                p["date"] = TODAY
            ep = path.replace("{t}", t)
            key = (prov, ep, json.dumps(p, sort_keys=True))
            if key not in cache:
                cache[key] = _http("UW" if prov == "UW" else "FH", ep, p)
            calls.append((t, cache[key]))
        errs = [(t, c) for t, c in calls if c["code"] is None or (c["code"] >= 400 and c["code"] != 404)]
        oks = [(t, c) for t, c in calls if c["code"] and c["code"] < 300 and not is_empty(c["data"])]
        codes = sorted({str(c["code"]) for _, c in calls if c["code"]})
        if errs:
            status = "ERROR"
            t0, c0 = errs[0]
            msg = f"{len(errs)}/{len(calls)} calls failed" + (f" (e.g. {t0}: {c0['error']})" if c0["error"] else "")
            code = ",".join(sorted({str(c["code"]) for _, c in errs if c["code"]})) or "n/a"
        elif not oks:
            status, msg = "NO DATA", "empty / null / zero for all tickers"
            code = ",".join(codes) or "n/a"
        else:
            status = "OK"
            msg = ("" if len(oks) == len(calls) else f"data for {len(oks)}/{len(calls)} tickers")
            code = ",".join(codes)
        smp = ""
        if oks:
            t0, c0 = oks[0]
            smp = (f"{t0}: " if t0 else "") + sample(c0["data"])
        rows.append({"Parameter": name, "Provider": "Unusual Whales" if prov == "UW" else "Finnhub",
                     "Endpoint": path + (("?" + "&".join(k for k in params)) if params else ""),
                     "Status": status, "HTTP code": code, "Error message": msg, "Sample value": smp,
                     "Used by": used, "Tickers OK": f"{len(oks)}/{len(calls)}"})
        print(f"  {status:8} {name}")
    rows += app_checks(tickers)
    for name, used, why in UNWIRED:
        rows.append({"Parameter": name, "Provider": "none", "Endpoint": "-", "Status": "NO DATA",
                     "HTTP code": "n/a", "Error message": why, "Sample value": "",
                     "Used by": used, "Tickers OK": "0/0"})
    write_xlsx(rows, out_path, tickers)
    return rows


def app_checks(tickers: list[str]) -> list[dict]:
    """Checks that go through the app's own code (what the scorers actually use):
    the VIX fallback value and the calendar-timing input."""
    out: list[dict] = []
    sys.path.insert(0, str(HERE))
    base = {"Provider": "Unusual Whales"}
    try:
        import kimi_service as ki
        lv = ki.vix_level()
        out.append({"Parameter": "VIX regime (app fallback value)", **base,
                    "Endpoint": "kimi_service.vix_level()", "Status": "OK" if lv else "NO DATA",
                    "HTTP code": "200" if lv else "n/a",
                    "Error message": "" if lv else "neither VIX index nor VIXY available",
                    "Sample value": f"{lv[0]:.2f} ({lv[1]})" if lv else "",
                    "Used by": "Kimi, Claude Pro", "Tickers OK": "1/1" if lv else "0/1"})
    except Exception as e:  # noqa: BLE001
        out.append({"Parameter": "VIX regime (app fallback value)", **base,
                    "Endpoint": "kimi_service.vix_level()", "Status": "ERROR", "HTTP code": "n/a",
                    "Error message": f"{type(e).__name__}: {e}"[:200], "Sample value": "",
                    "Used by": "Kimi, Claude Pro", "Tickers OK": "0/1"})
    try:
        import claude_last_service as cl
        import uw_earnings_calendar as cal
        rows = (cal.calendar(days=7) or {}).get("rows") or []
        timed = [r for r in rows if cl.timing_bias(r, 0.5) is not None]
        ok = bool(timed)
        out.append({"Parameter": "BMO/AMC calendar timing (scored)", **base,
                    "Endpoint": "/api/earnings/premarket|afterhours -> report_time",
                    "Status": "OK" if ok else "NO DATA", "HTTP code": "200" if ok else "n/a",
                    "Error message": "" if ok else "no row carries a BMO/AMC time",
                    "Sample value": (f"{timed[0].get('symbol')}: {timed[0].get('report_time')} -> "
                                     f"bias {cl.timing_bias(timed[0], 0.5)} ({len(timed)}/{len(rows)} rows timed)") if ok else "",
                    "Used by": "Claude, Grok, Meta, ChatGPT, Kimi", "Tickers OK": f"{len(timed)}/{len(rows)}"})
    except Exception as e:  # noqa: BLE001
        out.append({"Parameter": "BMO/AMC calendar timing (scored)", **base,
                    "Endpoint": "uw_earnings_calendar.calendar()", "Status": "ERROR", "HTTP code": "n/a",
                    "Error message": f"{type(e).__name__}: {e}"[:200], "Sample value": "",
                    "Used by": "Claude, Grok, Meta, ChatGPT, Kimi", "Tickers OK": "0/0"})
    return out


def write_xlsx(rows: list[dict], path: Path, tickers: list[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Health"
    cols = ["Parameter", "Provider", "Endpoint", "Status", "HTTP code", "Error message",
            "Sample value", "Used by", "Tickers OK"]
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    fills = {"OK": "C6EFCE", "NO DATA": "FFEB9C", "ERROR": "FFC7CE"}
    for row in ws.iter_rows(min_row=2):
        row[3].fill = PatternFill("solid", fgColor=fills.get(row[3].value, "FFFFFF"))
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHI", (42, 16, 46, 10, 11, 48, 60, 40, 11)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.append([])
    ws.append([f"Run {date.today().isoformat()} · tickers: {', '.join(tickers)}"])
    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(HERE.parent / "Data_Health_Check.xlsx"))
    a = ap.parse_args()
    rows = run(Path(a.out))
    bad = [r for r in rows if r["Status"] != "OK"]
    print(f"\nSaved {a.out}  ({len(rows)} parameters, {len(bad)} failing)\n")
    print("FAILING PARAMETERS")
    for r in bad:
        print(f"  [{r['Status']}] {r['Parameter']} — {r['Provider']} {r['Endpoint'].split('?')[0]} "
              f"(HTTP {r['HTTP code']}): {r['Error message']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
